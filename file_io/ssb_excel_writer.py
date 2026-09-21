"""Compact scientific Excel report for Engine-B SSB sessions."""

from __future__ import annotations

import math
import re
from dataclasses import dataclass
from datetime import datetime, timezone
from io import BytesIO
from typing import Iterable, Optional

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableColumn, TableStyleInfo
from openpyxl.worksheet.filters import AutoFilter
from domain.layers import cycle_data_equal

from config.constants import APP_VERSION
from domain.models import CycleData, ProcessingResult, Sample, UncertaintyBudget
from domain.ratio_selection import get_delta_cycle_data, get_ssb_cycle_data
from domain.uncertainty.scope import is_invalid_budget_scope
from file_io.report_formatting import report_uncertainty
from file_io.sanitize import format_contributor_name, format_isotope_label, safe_float, literal_excel_text

EXCEL_MAX_ROWS = 1_048_576
EXCEL_MAX_COLUMNS = 16_384


@dataclass(frozen=True)
class SSBReportOptions:
    include_uncertainty_detail: bool = True
    include_cycle_data: bool = False


def ssb_sheet_names(options: SSBReportOptions, *, has_detail: bool) -> list[str]:
    names = ["Summary", "Results"]
    if options.include_uncertainty_detail and has_detail:
        names.append("Budget_Detail")
    if options.include_cycle_data:
        names.append("Cycle_Data")
    return names


def ssb_has_detail(result: ProcessingResult) -> bool:
    """Whether at least one eligible Engine-B budget can populate detail."""
    return any(
        _valid_budget(sample, ratio) is not None
        for sample in result.samples
        if sample.is_sample
        for ratio in (sample.uncertainty or {})
    )


def _ratio_names(result: ProcessingResult) -> list[str]:
    names: set[str] = set()
    for sample in result.samples:
        for mapping in (sample.ratios, sample.blank_corrected_ratios, sample.corrected_ratios, sample.ssb_results,
                        sample.delta_results, sample.uncertainty):
            names.update(mapping or {})
    return sorted(names)


def _finite(value) -> Optional[float]:
    value = safe_float(value)
    return value if value is not None and math.isfinite(value) else None


def _stats(cd: Optional[CycleData]) -> tuple[Optional[float], ...]:
    if cd is None or not cd.n_valid:
        return (None, None, None, None, None, None)
    return (cd.mean, _finite(cd.sd), _finite(cd.se),
            _finite(2.0 * cd.se), _finite(cd.rsd_percent), cd.n_valid)


def _blank_corrected(sample: Sample, ratio: str) -> Optional[CycleData]:
    if sample.correction_records.get("hg") and ratio not in sample.blank_corrected_ratios:
        # A historical Hg-corrected working layer is not evidence of a blank snapshot.
        return None
    return (sample.blank_corrected_ratios or {}).get(ratio) or (sample.corrected_ratios or {}).get(ratio)


def _valid_budget(sample: Sample, ratio: str) -> Optional[UncertaintyBudget]:
    from domain.pb_correction_records import hg_blocks_final_value
    if hg_blocks_final_value(sample, ratio):
        return None
    budget = (sample.uncertainty or {}).get(ratio)
    if budget is None or is_invalid_budget_scope(budget):
        return None
    if (budget.engine or "") not in ("", "ssb_delta"):
        return None
    return budget


def _uncertainty_values(budget: Optional[UncertaintyBudget], ssb_mean: Optional[float]):
    if budget is None:
        return (None,) * 6
    u_rel = _finite(budget.u_combined_rel_permil)
    U_rel = _finite(budget.expanded_rel_permil)
    basis = _finite(budget.basis_ratio_value)
    reference = _finite(budget.delta_reference_value)
    if basis is None or basis == 0 or ssb_mean is None:
        return (None,) * 6
    # The stored basis must describe the displayed SSB layer. This strict check
    # prevents a final drift/calibrated budget being attached to an earlier stage.
    if not math.isclose(basis, ssb_mean, rel_tol=1e-10, abs_tol=1e-15):
        return (None,) * 6
    u_ratio = abs(basis) * u_rel / 1000.0 if u_rel is not None else None
    U_ratio = abs(basis) * U_rel / 1000.0 if U_rel is not None else None
    if reference is None or reference <= 0:
        return (u_rel, U_rel, u_ratio, U_ratio, None, None)
    scale = abs(basis / reference)
    return (u_rel, U_rel, u_ratio, U_ratio,
            scale * u_rel if u_rel is not None else None,
            scale * U_rel if U_rel is not None else None)


def _header(ws, row: int, headers: Iterable[object]) -> None:
    for col, value in enumerate(headers, 1):
        cell = ws.cell(row, col, value)
        cell.font = Font(
            bold=True,
            italic=isinstance(value, str) and value in {"SSB K", "K", "k"},
            color="FFFFFF",
        )
        cell.fill = PatternFill("solid", fgColor="2E75B6")
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")


def _delta_heading(ratio_label: str, *, prefix: str = "", suffix: str = " (‰)"):
    """Return native rich text with italic delta and upright isotope label."""
    ratio_label = re.sub(r"[A-Za-z]+/", "/", ratio_label)
    return CellRichText(
        *[TextBlock(InlineFont(b=True, color="FFFFFFFF", i=italic), text)
          for text, italic in [(prefix, False), ("δ", True), (ratio_label, False), (suffix, False)] if text],
    )


def _correction_heading():
    return CellRichText(TextBlock(InlineFont(b=True, i=True, color="FFFFFFFF"), "K"),
                        TextBlock(InlineFont(b=True, color="FFFFFFFF"), " factor"))


def _table(ws, header_row, last_row, headers):
    if last_row <= header_row:
        return
    name = f"{ws.title.replace('_', '')}_{header_row}"
    names = [str(h) for h in headers]
    if len({h.casefold() for h in names}) != len(names):
        raise ValueError(f"{ws.title}: Excel table headings must be unique ignoring case")
    table = Table(displayName=name, ref=f"A{header_row}:{get_column_letter(len(headers))}{last_row}")
    table.tableColumns = [TableColumn(id=i, name=str(h)) for i, h in enumerate(headers, 1)]
    table.autoFilter = AutoFilter(ref=table.ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)


def _pair(value_cell, uncertainty_cell):
    u = uncertainty_cell.value
    if isinstance(u, (float, int)) and u > 0:
        decimals = max(0, 1 - math.floor(math.log10(u)))
        fmt = "0" + ("." + "0" * decimals if decimals else "")
        value_cell.number_format = uncertainty_cell.number_format = fmt


def _finish(ws) -> None:
    ws.sheet_view.showGridLines = False
    for column in range(1, ws.max_column + 1):
        values = [str(ws.cell(row, column).value or "") for row in range(1, min(ws.max_row, 80) + 1)]
        if ws.title != "Summary":
            ws.column_dimensions[get_column_letter(column)].width = min(32, max(12, max(map(len, values), default=8) + 2))
    for cells in ws:
        for cell in cells:
            cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
            if isinstance(cell.value, float) and cell.number_format == "General":
                # Decimal display for small SE/u values; preserve the stored number.
                magnitude = abs(cell.value)
                places = max(0, 3 - math.floor(math.log10(magnitude))) if magnitude else 3
                cell.number_format = "0." + "0" * places if places else "0"
        ws.row_dimensions[cells[0].row].height = 60 if any(c.fill.fgColor.rgb == "002E75B6" for c in cells) else 32
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = False


def _summary(wb, result, processing_config, uncertainty_config, loaded_filename):
    ws = wb.active
    ws.title = "Summary"
    has_ssb = any(s.ssb_results for s in result.samples)
    has_delta = any(s.delta_results for s in result.samples)
    rows = [
        ("TraceISO Scientific Report", None),
        ("Source filename", loaded_filename),
        ("Software version", APP_VERSION),
        ("Exported UTC", datetime.now(timezone.utc).isoformat()),
        ("Element", result.element_symbol),
        ("Workflow", "Standard-sample bracketing" if has_ssb else "Raw and blank-corrected ratios"),
        ("Primary reporting mode", getattr(uncertainty_config, "output_mode", "delta") if has_delta else "absolute_ratio"),
        ("Reference material", getattr(processing_config, "reference_material", None)),
        ("Coverage method", {"fixed_k": "Fixed coverage factor", "welch_satterthwaite": "Welch–Satterthwaite"}.get(
            getattr(uncertainty_config, "coverage_method", None), getattr(uncertainty_config, "coverage_method", None))),
        ("Coverage factor k", getattr(uncertainty_config, "coverage_k", None)
         if getattr(uncertainty_config, "coverage_method", "fixed_k") == "fixed_k" else "Per result"),
        ("Blank method", getattr(processing_config, "blank_mode", None)),
        ("Outlier method", getattr(processing_config, "filter_method", None)),
        ("Outlier threshold", getattr(processing_config, "get_active_filter_threshold", lambda: None)()),
        ("SSB mode", getattr(processing_config, "ssb_mode", None)),
        ("Drift", "Enabled" if getattr(getattr(processing_config, "drift", None), "enabled", False) else "Disabled"),
        ("Observations", len(result.samples)),
    ]
    if not has_ssb:
        rows = [r for r in rows if r[0] not in {"SSB mode", "Reference material"}]
    if result.element_symbol == "Pb":
        from domain.pb_correction_records import hg_records
        rows.append(("Hg correction requested", bool(getattr(processing_config, "apply_hg_interference_correction", False))))
        evidence = sorted({f"{r.status}: {r.source or 'source unavailable'}; {r.intensity_basis}; {r.reason}"
                           for s in result.samples for r in hg_records(s).values()})
        rows += [("Hg correction evidence", text) for text in evidence]
    for row in rows:
        ws.append([literal_excel_text(v) if isinstance(v, str) else v for v in row])
    from config.reference_materials import get_all_certified_ratios, standard_uncertainty_from_values
    references = get_all_certified_ratios(result.element_symbol, crm_name=getattr(processing_config, "reference_material", None))
    if not (has_ssb or has_delta):
        references = {}
    for ratio, reference in references.items():
        ws.append((f"Reference ratio {format_isotope_label(ratio)}", reference.ratio))
        ws.cell(ws.max_row, 2).number_format = "0.###############"
        ws.append((f"Reference u ({format_isotope_label(ratio)}, k=1)",
                   standard_uncertainty_from_values(reference.uncertainty, reference.k)))
    for ratio in _ratio_names(result):
        if (has_ssb or has_delta) and ratio not in references:
            values = sorted({value for sample in result.samples
                             if (b := _valid_budget(sample, ratio)) is not None
                             and (value := _finite(b.certified_reference_value)) is not None})
            for value in values:
                ws.append((f"Reference ratio {format_isotope_label(ratio)}", value))
                ws.cell(ws.max_row, 2).number_format = "0.###############"
    ws["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    ws["A1"].alignment = Alignment(vertical="center")
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 90
    ws.freeze_panes = "A2"
    _finish(ws)


def _results(wb, result):
    ws = wb.create_sheet("Results")
    row = 1
    for ratio in _ratio_names(result):
        has_ssb = any(ratio in s.ssb_results for s in result.samples)
        has_delta = any(ratio in s.delta_results for s in result.samples)
        has_drift = any(ratio in s.drift_corrected_ratios for s in result.samples)
        from domain.pb_correction_records import hg_records, hg_blocks_final_value
        has_hg = result.element_symbol == "Pb" and any(
            ratio in s.interference_corrected_ratios or any(r.requested and r.ratio_name == ratio for r in hg_records(s).values())
            for s in result.samples)
        label = format_isotope_label(ratio)
        ws.cell(row, 1, f"Ratio: {label}").font = Font(bold=True, color="FFFFFF")
        ws.cell(row, 1).fill = PatternFill("solid", fgColor="1F4E78")
        row += 1
        headers = ["Sample", "Type", "Run", "Accepted n", f"Blank-corrected {label}", "SD", "SE", "2SE", "RSD (%)",
                   f"SSB-corrected {label}", f"SSB-corrected {label} SD", f"SSB-corrected {label} SE", _correction_heading(), "Relative u_c(R_SSB) (‰)",
                   "Relative U(R_SSB) (‰)", "U(R_SSB)", _delta_heading(label),
                   _delta_heading(label, prefix="SD(", suffix=") (‰)"),
                   _delta_heading(label, prefix="SE(", suffix=") (‰)"),
                   _delta_heading(label, prefix="2SE(", suffix=") (‰)"),
                   _delta_heading(label, prefix="U(", suffix=") (‰)"), "k", "Dominant uncertainty contributor"]
        if has_drift:
            headers += [f"Drift-adjusted {label}", "Drift SD", "Drift SE",
                        "Relative u_c(R_drift) (‰)", "Relative U(R_drift) (‰)", "U(R_drift)"]
        # Select stage columns before writing, preserving their matching values.
        columns = list(range(len(headers)))
        if not has_delta:
            columns = [i for i in columns if i not in range(16, 21)]
        if not has_ssb:
            columns = [i for i in columns if i not in range(9, 16)]
            if not has_delta and not has_drift:
                columns = [i for i in columns if i not in (21, 22)]
            raw_start = len(headers)
            headers += [f"Raw {label}", "Raw SD", "Raw SE", "Raw 2SE", "Raw RSD (%)"]
            columns[4:4] = list(range(raw_start, raw_start + 5))
        hg_start = len(headers)
        if has_hg:
            headers += [f"Hg-corrected {label}", "Hg SD", "Hg SE", "Hg correction status"]
            insert = next((i for i, value in enumerate(columns) if value >= 9 and value < 23), len(columns))
            columns[insert:insert] = list(range(hg_start, hg_start + 4))
        headers = [headers[i] for i in columns]
        _header(ws, row, headers)
        header_row = row
        row += 1
        for sample in result.samples:
            if sample.is_blank:
                continue
            bc = _stats(_blank_corrected(sample, ratio))
            ssb_cd = get_ssb_cycle_data(sample, ratio)
            ssb = _stats(ssb_cd)
            delta = (sample.delta_results or {}).get(ratio, {}) or {}
            blocked_hg = result.element_symbol == "Pb" and hg_blocks_final_value(sample, ratio)
            if blocked_hg:
                ssb_cd, ssb, delta = None, _stats(None), {}
            budget = _valid_budget(sample, ratio)
            uv = _uncertainty_values(budget, ssb[0])
            delta_uv = _uncertainty_values(budget, _finite(delta.get("sample_mean", ssb[0])))
            if budget and delta.get("std_mean") is not None and delta.get("std_mean") != budget.delta_reference_value:
                delta_uv = (None,) * 6
            values = [literal_excel_text(sample.name), literal_excel_text(sample.sample_type), sample.run_number,
                      ssb[5] if ssb_cd is not None else (bc[5] if bc[5] is not None else _stats(sample.ratios.get(ratio))[5]), *bc[:5], *ssb[:3],
                      _finite(((sample.ssb_results or {}).get(ratio, {}) or {}).get("k_factor")),
                      uv[0], uv[1], report_uncertainty(uv[3]) if uv[3] is not None else None,
                      _finite(delta.get("delta")), _finite(delta.get("delta_sd")), _finite(delta.get("delta_se")),
                      _finite(delta.get("delta_2se")), report_uncertainty(delta_uv[5]) if delta_uv[5] is not None and delta else None,
                      _finite(budget.coverage_factor_k) if budget else None,
                      literal_excel_text(format_contributor_name(budget.dominant_contributor)) if budget else (
                          "Not quantified" if ssb_cd is not None and sample.is_sample else None
                      )]
            if budget and ssb_cd is not None and uv[3] is None:
                values[15] = "Not quantified for SSB layer"
            if has_drift:
                drift = _stats(None if blocked_hg else sample.drift_corrected_ratios.get(ratio))
                drift_uv = _uncertainty_values(budget, drift[0])
                values += [*drift[:3], drift_uv[0], drift_uv[1],
                           report_uncertainty(drift_uv[3]) if drift_uv[3] is not None else None]
            if not has_ssb:
                values += list(_stats(sample.ratios.get(ratio))[:5])
            if has_hg:
                record = hg_records(sample).get(ratio)
                values += list(_stats(sample.interference_corrected_ratios.get(ratio))[:3]) + [
                    literal_excel_text(f"{record.status}: {record.reason}" if record else "Not requested for this ratio")]
            ws.append([values[i] for i in columns])
            for cell in ws[row]:
                if isinstance(cell.value, float):
                    magnitude = abs(cell.value)
                    decimals = max(0, 3 - math.floor(math.log10(magnitude))) if magnitude else 3
                    cell.number_format = "0." + "0" * min(decimals, 12) if decimals else "0"
            positions = {original: col for col, original in enumerate(columns, 1)}
            for value_index, uncertainty_index in ((9, 15), (16, 20), (23, 28)):
                if value_index in positions and uncertainty_index in positions:
                    _pair(ws.cell(row, positions[value_index]), ws.cell(row, positions[uncertainty_index]))
            if 21 in positions:
                ws.cell(row, positions[21]).number_format = "0.###"
            row += 1
        _table(ws, header_row, row - 1, headers)
        row += 2
    ws.freeze_panes = "E3"
    ws.print_title_rows = "1:2"
    _finish(ws)


def _contributor_names(rows) -> list[str]:
    names: list[str] = []
    for _sample, budget in rows:
        for contributor in budget.contributors or []:
            if (contributor.is_active or contributor.state in {"MISSING_DATA", "NO_APPROVED_MODEL"}) and contributor.name not in names:
                names.append(contributor.name)
    return names


def _budget_detail(wb, result):
    ws = wb.create_sheet("Budget_Detail")
    row = 1
    for ratio in _ratio_names(result):
        rows = [(s, b) for s in result.samples if s.is_sample and (b := _valid_budget(s, ratio)) is not None]
        if not rows:
            continue
        has_ssb = any(ratio in s.ssb_results for s, _ in rows)
        has_delta = any(ratio in s.delta_results for s, _ in rows)
        contributors = _contributor_names(rows)
        ws.cell(row, 1, f"Ratio: {format_isotope_label(ratio)}").font = Font(bold=True, color="FFFFFF")
        ws.cell(row, 1).fill = PatternFill("solid", fgColor="1F4E78")
        row += 1
        headers = ["Sample", "Run", "Preceding standard mean", "Preceding standard SE", "Following standard mean",
                   "Following standard SE", "Bracketing mean", "R_SSB", "SE(R_SSB)"]
        headers += [f"{format_contributor_name(n)} (‰)" for n in contributors]
        headers += ["Relative u_c(R_SSB) (‰)", "Relative U(R_SSB) (‰)", "k", "u_c(R_SSB)", "U(R_SSB)", "u_c(δ) (‰)", "U(δ) (‰)"]
        if not has_delta:
            headers = headers[:-2]
        if not has_ssb:
            headers = headers[:2] + headers[7:]
            headers = [h.replace("R_SSB", "R_blank") for h in headers]
        _header(ws, row, headers)
        header_row = row
        row += 1
        for sample, budget in rows:
            payload = ((sample.ssb_results or {}).get(ratio, {}) or {})
            ssb = _stats(get_ssb_cycle_data(sample, ratio) if has_ssb else _blank_corrected(sample, ratio))
            uv = _uncertainty_values(budget, ssb[0])
            delta = sample.delta_results.get(ratio, {})
            delta_uv = _uncertainty_values(budget, _finite(delta.get("sample_mean", ssb[0])))
            if delta.get("std_mean") is not None and delta.get("std_mean") != budget.delta_reference_value:
                delta_uv = (None,) * 6
            by_name = {c.name: c for c in budget.contributors or []}
            cvalues = []
            for name in contributors:
                c = by_name.get(name)
                if c is None or c.state in {"BY_SAMPLE_DESIGN", "BY_GLOBAL_DESIGN", "NOT_APPLICABLE"}:
                    cvalues.append(None)
                elif c.state in {"MISSING_DATA", "NO_APPROVED_MODEL"}:
                    cvalues.append("Not quantified")
                else:
                    cvalues.append(_finite(c.value_rel_permil))
            prev, nxt = _finite(payload.get("prev_std_mean")), _finite(payload.get("next_std_mean"))
            bracket = (prev + nxt) / 2 if prev is not None and nxt is not None else _finite(payload.get("bracketing_avg"))
            values = [literal_excel_text(sample.name), sample.run_number, prev, _finite(payload.get("prev_std_se")),
                       nxt, _finite(payload.get("next_std_se")), bracket, ssb[0], ssb[2], *cvalues,
                       uv[0], uv[1], _finite(budget.coverage_factor_k), uv[2], uv[3],
                       delta_uv[4] if delta else None, delta_uv[5] if delta else None]
            if not has_delta:
                values = values[:-2]
            if not has_ssb:
                values = values[:2] + values[7:]
            ws.append(values)
            row += 1
        _table(ws, header_row, row - 1, headers)
        row += 2
    ws.freeze_panes = "B3"
    _finish(ws)


def _series(sample: Sample):
    out: list[tuple[str, str, str, CycleData]] = []
    mappings = [
        ("Raw intensity", "V", sample.intensities), ("Blank-corrected intensity", "V", sample.blank_corrected_intensities),
        ("Corrected intensity", "V", sample.corrected_intensities), ("Raw ratio", "1", sample.ratios),
        ("Blank-corrected ratio", "1", sample.blank_corrected_ratios), ("Corrected ratio", "1", sample.corrected_ratios),
        ("Drift-adjusted ratio", "1", sample.drift_corrected_ratios),
    ]
    if sample.interference_corrected_ratios or sample.correction_records.get("hg"):
        mappings[2:2] = [("Hg-corrected intensity", "V", sample.interference_corrected_intensities)]
        mappings.insert(-1, ("Hg-corrected ratio", "1", sample.interference_corrected_ratios))
    for stage, unit, mapping in mappings:
        for name, cd in (mapping or {}).items():
            if cd is not None:
                if stage.startswith("Corrected") and any(
                    prev_stage.startswith("Blank-corrected") and prev_name == name and prev_unit == unit
                    and cycle_data_equal(prev_cd, cd)
                    for prev_stage, prev_name, prev_unit, prev_cd in out
                ):
                    continue
                out.append((stage, name, unit, cd))
    for ratio in sorted(set((sample.ssb_results or {})) | set((sample.delta_results or {}))):
        ssb, delta = get_ssb_cycle_data(sample, ratio), get_delta_cycle_data(sample, ratio)
        if ssb is not None: out.append(("SSB-corrected ratio", ratio, "1", ssb))
        if delta is not None: out.append(("Delta", ratio, "‰", delta))
    return out


def _cycle_data(wb, result):
    specifications: list[tuple[str, str, str]] = []
    per_sample = []
    total_rows = 1
    for sample in result.samples:
        series = _series(sample)
        per_sample.append((sample, series))
        total_rows += max((cd.n_total for *_key, cd in series), default=sample.n_cycles)
        for stage, name, unit, _cd in series:
            key = (stage, name, unit)
            if key not in specifications: specifications.append(key)
    column_count = 4 + 2 * len(specifications)
    if total_rows > EXCEL_MAX_ROWS or column_count > EXCEL_MAX_COLUMNS:
        raise ValueError(f"Cycle_Data dimensions {total_rows} rows x {column_count} columns exceed Excel limits")
    ws = wb.create_sheet("Cycle_Data")
    headers = ["Sample", "Type", "Run", "Cycle"]
    for stage, name, unit in specifications:
        headers.extend((f"{stage}: {format_isotope_label(name)} ({unit})", f"{stage}: {format_isotope_label(name)} status"))
    _header(ws, 1, headers)
    duplicates = {}
    for sample in result.samples:
        duplicates[(sample.name, sample.run_number)] = duplicates.get((sample.name, sample.run_number), 0) + 1
    row = 2
    for sample, series in per_sample:
        by_key = {(a, b, c): cd for a, b, c, cd in series}
        count = max((cd.n_total for cd in by_key.values()), default=sample.n_cycles)
        name = sample.name
        if duplicates[(sample.name, sample.run_number)] > 1:
            name = f"{name} [{sample.observation_id}]"
        for index in range(count):
            values = [literal_excel_text(name), literal_excel_text(sample.sample_type), sample.run_number, index + 1]
            for key in specifications:
                cd = by_key.get(key)
                if cd is None or index >= cd.n_total:
                    values.extend((None, "Not recorded"))
                else:
                    values.extend((_finite(cd.values[index]), "Accepted" if bool(cd.mask[index]) else "Excluded"))
            ws.append(values); row += 1
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:{get_column_letter(column_count)}{row - 1}"
    ws.print_title_rows = "1:1"
    _finish(ws)


def export_ssb_report(result: ProcessingResult, *, options: SSBReportOptions,
                      cycle_result: Optional[ProcessingResult] = None,
                      processing_config=None, uncertainty_config=None,
                      loaded_filename: Optional[str] = None, provenance=None, include_metadata_sheets=True) -> BytesIO:
    wb = Workbook()
    wb._traceiso_include_metadata_sheets = include_metadata_sheets
    _summary(wb, result, processing_config, uncertainty_config, loaded_filename)
    _results(wb, result)
    has_detail = ssb_has_detail(result)
    if options.include_uncertainty_detail and has_detail:
        _budget_detail(wb, result)
    if options.include_cycle_data:
        _cycle_data(wb, cycle_result or result)
    expected = ssb_sheet_names(options, has_detail=has_detail)
    if wb.sheetnames != expected:
        raise ValueError(f"SSB workbook sheet plan mismatch: expected {expected}, wrote {wb.sheetnames}")
    output = BytesIO()
    from file_io.excel_writer import finalize_workbook_evidence
    finalize_workbook_evidence(wb, result, provenance=provenance, uncertainty_config=uncertainty_config)
    wb.save(output)
    output.seek(0)
    return output
