"""Compact Pb–Tl reports of supplied layers and Engine-C budgets.

This boundary validates producer evidence. It never replays a correction or
claims historical mask identity from matching means and counts.
"""

from collections import Counter
from dataclasses import dataclass
from io import BytesIO
import math

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from domain.layers import cycle_data_equal
from domain.output_scale import resolve_output_scale
from domain.pb_calibration_records import calibration_records, calibrated_delta_record, TARGET_ROLES
from domain.pb_correction_records import hg_records
from domain.ratio_selection import select_best_ratio_layer
from domain.uncertainty.scope import is_invalid_budget_scope, budget_scope_note
from file_io.report_formatting import report_uncertainty
from file_io.sanitize import format_isotope_label
from file_io.sr_excel_writer import (
    _append, _finite, _finish, _header, _stats, _uncertainty_headers,
    _uncertainty_values,
)
from file_io.ssb_excel_writer import _pair, _table

EXCEL_MAX_ROWS = 1_048_576
EXCEL_MAX_COLUMNS = 16_384


@dataclass(frozen=True)
class PbReportOptions:
    include_uncertainty_detail: bool = True
    include_cycle_data: bool = False


def resolve_pb_options(options=None, *, detail=True, cycles=False):
    if options is None:
        return PbReportOptions(detail, cycles)
    if not isinstance(options, PbReportOptions):
        raise TypeError("pb_report_options must be a PbReportOptions instance")
    return options


def pb_sheet_names(options, *, has_detail):
    options = resolve_pb_options(options)
    return (["Summary", "Results"]
            + (["Budget_Detail"] if options.include_uncertainty_detail and has_detail else [])
            + (["Cycle_Data"] if options.include_cycle_data else []))


def _ratio_names(result):
    return sorted({r for s in result.samples for attr in (
        "ratios", "blank_corrected_ratios", "interference_corrected_ratios",
        "corrected_ratios", "iif_corrected_ratios", "pb_standard_corrected_ratios",
        "drift_corrected_ratios", "uncertainty",
    ) for r in getattr(s, attr) if "Pb" in r})


def _calibration_requested(result, cfg=None):
    return bool(getattr(getattr(cfg, "pb_standard_calibration", None), "enabled", False)
                or any(calibration_records(s) for s in result.samples))


def validate_pb_report(result, processing_config=None, uncertainty_config=None, calibration_freshness=None):
    if result.element_symbol != "Pb":
        raise ValueError("The Pb–Tl report requires a Pb session")
    if calibration_freshness is not None and calibration_freshness.get("status") == "stale":
        raise ValueError("Pb-standard calibration is stale; rerun data reduction")
    if processing_config is not None:
        if not processing_config.apply_mass_bias_correction:
            raise ValueError("The Pb–Tl report requires external Tl normalization; reprocess or use the SSB report")
        if uncertainty_config is not None and uncertainty_config.resolve_engine(
            "Pb", processing_config=processing_config
        ) != "pb_tl_external_normalization":
            raise ValueError("Incompatible Pb–Tl uncertainty route")
    if _calibration_requested(result, processing_config):
        for sample in result.samples:
            if sample.is_blank:
                continue
            for ratio in _ratio_names(result):
                if any(ratio in getattr(sample, a) for a in (
                    "ratios", "corrected_ratios", "iif_corrected_ratios", "pb_standard_corrected_ratios"
                )) and ratio not in calibration_records(sample):
                    raise ValueError(f"Requested Pb calibration has no record for {sample.name}, {ratio}; reprocess")


def pb_report_context(sample, ratio, *, processing_config=None, calibration_requested=False):
    """One value/uncertainty eligibility rule for Results, preview and detail."""
    record = calibration_records(sample).get(ratio)
    if sample.metadata.get("excluded", False):
        return None, None, "Observation excluded"
    if record is not None:
        if record.role not in TARGET_ROLES:
            return None, None, record.reason or "Not applicable to this calibration role"
        if record.status != "applied":
            return None, None, record.reason or "Calibration unavailable"
        if ratio not in sample.pb_standard_corrected_ratios:
            return None, None, "Recorded calibrated layer missing"
    elif calibration_requested:
        return None, None, "Requested calibration record missing"
    if ratio not in sample.iif_corrected_ratios:
        return None, None, "Tl-normalized layer unavailable"
    layer = select_best_ratio_layer(sample, ratio)
    expected_layer = "pb_standard" if record else "iif"
    drift = getattr(processing_config, "drift", None)
    drift_requested = not record and bool(getattr(drift, "enabled", False)) and drift.ratio_name == ratio
    if drift_requested and ratio not in sample.drift_corrected_ratios:
        return None, None, "Requested drift layer unavailable"
    if layer is not None and layer.key == "drift" and not record:
        scale = resolve_output_scale(sample, ratio, input_layer="iif")
        if not any(name == "drift" for name, _ in scale.components) or (drift is not None and not drift_requested):
            return None, None, "Drift layer lacks current applied evidence"
        expected_layer = "drift"
    if layer is None or layer.key != expected_layer:
        return None, None, "Governing final Pb layer unavailable"
    mean = _finite(layer.data.mean)
    if mean is None or mean <= 0:
        return None, None, "No finite positive final ratio"
    budget = sample.uncertainty.get(ratio)
    if budget is None:
        return layer, None, "Budget unavailable"
    if is_invalid_budget_scope(budget) or budget.budget_scope == "not_calculated":
        return layer, None, budget_scope_note(budget) or "Not calculated"
    if budget.budget_scope not in {"stored", "hybrid", "full_runtime", "limited"}:
        return layer, None, "Unsupported budget scope"
    engine = "pb_tl_standard_calibration" if record else "pb_tl_external_normalization"
    if budget.engine != engine or budget.output_mode != "absolute_ratio":
        return layer, None, "Budget engine or result space mismatch"
    if layer.data.n_valid < 2:
        return layer, None, "Insufficient accepted cycles"
    # Both Engine-C producers combine on the final absolute scale. Their basis
    # is already calibrated/drift-adjusted; applying an output factor again
    # would double scale the uncertainty.
    if any(_finite(v) is None or not math.isclose(v, mean, rel_tol=1e-12, abs_tol=0)
           for v in (budget.ratio_value, budget.basis_ratio_value)):
        return layer, None, "Budget basis does not describe the reported ratio"
    if budget.n_cycles != layer.data.n_valid:
        return layer, None, "Budget accepted-cycle count mismatch"
    numbers = (budget.u_combined_abs, budget.expanded_abs, budget.u_combined_rel_permil,
               budget.expanded_rel_permil, budget.coverage_factor_k)
    if any(_finite(v) is None or v < 0 for v in numbers) or budget.coverage_factor_k <= 0:
        return layer, None, "Invalid uncertainty values"
    if budget.effective_dof != math.inf and (_finite(budget.effective_dof) is None or budget.effective_dof <= 0):
        return layer, None, "Invalid effective degrees of freedom"
    if any(c.state == "MISSING_DATA" for c in budget.contributors):
        return layer, None, "Contributor missing data"
    for c in budget.contributors:
        if c.is_active and any(_finite(v) is None or v < 0 for v in (c.value_abs, c.value_rel_permil)):
            return layer, None, "Invalid active contributor"
    for actual, expected in (
        (budget.expanded_abs, budget.u_combined_abs * budget.coverage_factor_k),
        (budget.u_combined_rel_permil, budget.u_combined_abs / mean * 1000),
        (budget.expanded_rel_permil, budget.expanded_abs / mean * 1000),
    ):
        if not math.isclose(actual, expected, rel_tol=1e-10, abs_tol=0):
            return layer, None, "Inconsistent uncertainty units or coverage"
    limitations = any(c.state == "NO_APPROVED_MODEL" for c in budget.contributors)
    return layer, budget, "Quantified terms only; no approved model for effects listed in Summary" if limitations else ""


def _contexts(result, cfg=None):
    calibrated = _calibration_requested(result, cfg)
    return {(s.observation_id, r): pb_report_context(s, r, processing_config=cfg,
            calibration_requested=calibrated) for s in result.samples if not s.is_blank
            for r in _ratio_names(result)}


def pb_has_detail(result, *, processing_config=None):
    return any(b is not None for _layer, b, _reason in _contexts(result, processing_config).values())


def _stages(sample, ratio, *, cfg=None, calibrated=False):
    stages = {}
    for key, label, mapping in (
        ("raw", "Raw", sample.ratios),
        ("blank", "Blank-corrected", sample.blank_corrected_ratios),
        ("hg", "Hg-corrected", sample.interference_corrected_ratios),
        ("tl", "Tl-normalized", sample.iif_corrected_ratios),
    ):
        if ratio in mapping:
            stages[key] = (label, mapping[ratio])
    record = calibration_records(sample).get(ratio)
    if record and record.role in TARGET_ROLES and record.status == "applied":
        cd = sample.pb_standard_corrected_ratios.get(ratio)
        if cd is not None:
            stages["calibration"] = (record.final_layer_label, cd)
    if not calibrated and ratio in sample.drift_corrected_ratios:
        scale = resolve_output_scale(sample, ratio, input_layer="iif")
        drift = getattr(cfg, "drift", None)
        if (scale.output_layer == "drift" and any(name == "drift" for name, _ in scale.components)
                and (drift is None or drift.enabled and drift.ratio_name == ratio)
                and ratio in sample.iif_corrected_ratios):
            stages["drift"] = ("Drift-corrected", sample.drift_corrected_ratios[ratio])
    return stages


def _identity(sample, duplicates):
    return sample.name + (f" [{sample.observation_id}]" if duplicates[(sample.name, sample.run_number)] > 1 else "")


def _delta_requested(result, cfg):
    if cfg is not None:
        from domain.pb_standard_calibration import calibrated_delta_requested
        return calibrated_delta_requested("Pb", cfg)
    return any(s.correction_records.get("pb_calibrated_delta") for s in result.samples)


def _delta_available(sample, ratio, layer, record, delta_record):
    if (delta_record is None or delta_record.status != "applied" or layer is None
            or record is None or delta_record.calibration_record_sha256 != record.sha256
            or delta_record.result_space != "delta_permil" or delta_record.n_valid < 1):
        return False
    payload = sample.delta_results.get(ratio, {})
    if payload:
        if payload.get("source_layer") != "pb_standard":
            return False
        n, mean = payload.get("n"), payload.get("delta")
    else:
        cd = sample.pb_calibrated_delta_cycles.get(ratio)
        n, mean = (cd.n_valid, cd.mean) if cd is not None else (None, None)
    return (n == delta_record.n_valid and _finite(mean) is not None
            and _finite(delta_record.delta_mean) is not None
            and math.isclose(mean, delta_record.delta_mean, rel_tol=1e-12, abs_tol=1e-12))


def _results(wb, result, cfg, uc, contexts):
    ws = wb.create_sheet("Results")
    observations = [s for s in result.samples if not s.is_blank]
    duplicates = Counter((s.name, s.run_number) for s in observations)
    calibrated = _calibration_requested(result, cfg)
    delta = _delta_requested(result, cfg)
    for ratio in _ratio_names(result):
        stages = [_stages(s, ratio, cfg=cfg, calibrated=calibrated) for s in observations]
        schema = [(k, label) for k, label in (
            ("raw", "Raw"), ("blank", "Blank-corrected"), ("hg", "Hg-corrected"),
            ("tl", "Tl-normalized"), ("calibration", "Pb-standard calibration after Tl"),
            ("drift", "Drift-corrected"),
        ) if any(k in row for row in stages) or k == "tl" or k == "calibration" and calibrated]
        _append(ws, [format_isotope_label(ratio)])
        header_row = ws.max_row + 1
        headers = ["Sample", "Type", "Run"] + (["Calibration role", "Applied calibration mode"] if calibrated else [])
        for key, label in schema:
            headers += [f"{label} ratio", f"{label} SD (ratio)", f"{label} SE (ratio)", f"{label} n"]
        value_col = len(headers) + 2
        headers += ["Final source layer", "Reported ratio", "Accepted n"] + _uncertainty_headers() + ["Availability"]
        if cfg is not None:
            requested_statistics = {cfg.pb_standard_calibration.delta_precision_statistic}
        else:
            # Historical observations may have different recorded selections.
            # Keep each statistic in its own column rather than relabeling it.
            requested_statistics = {dr.precision_statistic for s in observations
                                    if (dr := calibrated_delta_record(s, ratio)) is not None}
        statistics = [name for name in ("sd", "se") if name in requested_statistics]
        if delta:
            headers += ["Calibrated delta (‰)", "Delta accepted n", "Delta reference"]
            for statistic in statistics:
                headers += ["Cycle scatter (SD) (‰)" if statistic == "sd" else "Precision of the mean (SE) (‰)"]
            headers += ["Delta combined uncertainty"]
        _header(ws, header_row, headers)
        for s, stage in zip(observations, stages):
            record = calibration_records(s).get(ratio)
            layer, budget, reason = contexts[s.observation_id, ratio]
            values = [_identity(s, duplicates), s.sample_type, s.run_number]
            if calibrated:
                values += [record.role if record else None, record.applied_mode if record else None]
            for key, _label in schema:
                cd = stage.get(key, (None, None))[1]
                values += _stats(cd) + [cd.n_valid if cd is not None else None]
            label = record.final_layer_label if record and layer else "Tl-normalized" if layer and layer.key == "iif" else layer.label if layer else None
            values += [label, _finite(layer.data.mean) if layer else None, layer.data.n_valid if layer else None]
            values += _uncertainty_values(budget, show_dof=getattr(uc, "coverage_method", "") != "fixed_k") + [reason or "Quantified"]
            if delta:
                dr = calibrated_delta_record(s, ratio)
                available = _delta_available(s, ratio, layer, record, dr)
                payload = s.delta_results.get(ratio, {})
                values += [payload.get("delta", dr.delta_mean) if available else None,
                           payload.get("n", dr.n_valid) if available else None,
                           dr.reference_material if available else None]
                for statistic in statistics:
                    precision = payload.get("delta_" + statistic, dr.precision_value) if available else None
                    values += [precision if available and dr.n_valid >= 2 and dr.precision_statistic == statistic else None]
                values += ["Not calculated" if available else (dr.reason or "Delta record does not match prepared result" if dr else "Delta unavailable")]
            _append(ws, values)
            for col, heading in enumerate(headers, 1):
                if str(heading).endswith(" ratio"):
                    ws.cell(ws.max_row, col).number_format = "0.00000000"
            _pair(ws.cell(ws.max_row, value_col), ws.cell(ws.max_row, value_col + 3))
        _table(ws, header_row, ws.max_row, headers)
        ws.append([])
    ws.freeze_panes = "F3" if calibrated else "D3"
    ws.print_title_rows = "1:2"
    _finish(ws)
    ws.print_title_cols = "A:E" if calibrated else "A:C"


def _detail(wb, result, uc, contexts):
    ws = wb.create_sheet("Budget_Detail")
    duplicates = Counter((s.name, s.run_number) for s in result.samples)
    for ratio in _ratio_names(result):
        rows = [(s, *contexts[s.observation_id, ratio]) for s in result.samples
                if not s.is_blank and contexts[s.observation_id, ratio][1] is not None]
        if not rows:
            continue
        columns = {}
        for s, layer, budget, reason in rows:
            for c in budget.contributors:
                columns.setdefault(c.name, f"{c.display_name or c.name} [{c.name}] (‰)")
        _append(ws, [format_isotope_label(ratio)])
        header_row = ws.max_row + 1
        headers = ["Sample", "Type", "Run", "Calibration role", "Final source layer", "Reported ratio", "Accepted n"]
        headers += list(columns.values()) + _uncertainty_headers() + ["Availability"]
        _header(ws, header_row, headers)
        for s, layer, budget, reason in rows:
            record = calibration_records(s).get(ratio)
            contributors = {c.name: c for c in budget.contributors}
            values = [_identity(s, duplicates), s.sample_type, s.run_number, record.role if record else None,
                      record.final_layer_label if record else "Tl-normalized" if layer.key == "iif" else layer.label,
                      layer.data.mean, layer.data.n_valid]
            for name in columns:
                c = contributors.get(name)
                values.append(report_uncertainty(c.value_rel_permil) if c and c.is_active else
                              "No approved model" if c and c.state == "NO_APPROVED_MODEL" else None)
            values += _uncertainty_values(budget, show_dof=getattr(uc, "coverage_method", "") != "fixed_k")
            values += [reason or "Quantified"]
            _append(ws, values)
            _pair(ws.cell(ws.max_row, 6), ws.cell(ws.max_row, 9 + len(columns)))
        _table(ws, header_row, ws.max_row, headers)
        ws.append([])
    ws.freeze_panes = "D3"
    ws.print_title_rows = "1:2"
    _finish(ws)


def _series(sample, cfg=None):
    series = []
    for stage, unit, mapping in (
        ("Raw intensity", "V", sample.intensities),
        ("Blank-corrected intensity", "V", sample.blank_corrected_intensities),
        ("Hg-corrected intensity", "V", sample.interference_corrected_intensities),
        ("Corrected intensity", "V", sample.corrected_intensities),
        ("Raw ratio", "ratio", sample.ratios),
        ("Blank-corrected ratio", "ratio", sample.blank_corrected_ratios),
        ("Hg-corrected ratio", "ratio", sample.interference_corrected_ratios),
        ("Corrected ratio", "ratio", sample.corrected_ratios),
        ("Tl-normalized ratio", "ratio", sample.iif_corrected_ratios),
        ("Pb-standard-calibrated ratio", "ratio", sample.pb_standard_corrected_ratios),
        ("Calibrated delta", "‰", sample.pb_calibrated_delta_cycles),
    ):
        for name, cd in sorted(mapping.items()):
            if stage.startswith("Corrected") and any(n == name and u == unit and cycle_data_equal(cd, prev)
                for label, n, u, prev in series if label.startswith("Blank-corrected")):
                continue
            series.append((stage, name, unit, cd))
    if not calibration_records(sample):
        for ratio in sample.drift_corrected_ratios:
            stage = _stages(sample, ratio, cfg=cfg)
            if "drift" in stage:
                series.append(("Drift-corrected ratio", ratio, "ratio", stage["drift"][1]))
    return series


def _cycle_data(wb, result, cfg=None):
    prepared = [(s, _series(s, cfg)) for s in result.samples]
    specs = list(dict.fromkeys((stage, name, unit) for s, series in prepared for stage, name, unit, cd in series))
    counts = [max((cd.n_total for *_, cd in series), default=0) for s, series in prepared]
    rows, cols = 1 + sum(counts), 5 + 2 * len(specs)
    if rows > EXCEL_MAX_ROWS or cols > EXCEL_MAX_COLUMNS:
        raise ValueError(f"Cycle_Data dimensions {rows} rows x {cols} columns exceed Excel limits; disable Cycle_Data")
    ws = wb.create_sheet("Cycle_Data")
    headers = ["Sample", "Type", "Run", "Observation ID", "Cycle"]
    for stage, name, unit in specs:
        title = f"{stage}: {format_isotope_label(name)} ({unit})"
        headers += [title, title + " accepted"]
    _header(ws, 1, headers)
    for (sample, series), count in zip(prepared, counts):
        lookup = {(stage, name, unit): cd for stage, name, unit, cd in series}
        for i in range(count):
            values = [sample.name, sample.sample_type, sample.run_number, sample.observation_id, i + 1]
            for key in specs:
                cd = lookup.get(key)
                values += ([_finite(cd.values[i]), bool(cd.mask[i]) and _finite(cd.values[i]) is not None]
                           if cd is not None and i < cd.n_total else [None, None])
            _append(ws, values)
    ws.freeze_panes = "F2"
    ws.print_title_rows = "1:1"
    ws.auto_filter.ref = f"A1:{get_column_letter(cols)}{rows}"
    _finish(ws)
    ws.column_dimensions["D"].width = 44
    # Formatting only; all archived cells retain their unrounded numeric value.
    for row in ws.iter_rows(min_row=2, min_col=6):
        for cell in row[::2]:
            cell.number_format = "0.0000000000E+00"


def _summary(wb, result, cfg, uc, filename, element, contexts):
    ws = wb.active
    ws.title = "Summary"
    _header(ws, 1, ["Pb–Tl scientific report", "Session context"])
    rows = [("Source filename", filename), ("Element", "Pb"),
            ("Workflow", "External Tl normalization"),
            ("Reported ratios", ", ".join(format_isotope_label(r) for r in _ratio_names(result))),
            ("Observations", len(result.samples)), ("Uncertainty units", "Absolute ratio and relative ‰; U = k × u_c"),
            ("Cross-ratio scope", "No cross-ratio covariance is propagated; do not treat ratios as independent in downstream combinations."),
            ("Coverage method", getattr(uc, "coverage_method", None))]
    if cfg is not None:
        rows += [("Blank method", cfg.blank_mode), ("Outlier method", cfg.filter_method),
                 ("Outlier threshold", cfg.get_active_filter_threshold()),
                 ("Normalization pair", format_isotope_label(cfg.normalization_ratio_override or getattr(element, "normalization_ratio", "") or "")),
                 ("Accepted Tl ratio", cfg.normalization_value_override if cfg.normalization_value_override is not None else getattr(element, "normalization_value", None)),
                 ("Correction law", "Russell exponential law"),
                 ("Hg correction requested", "Yes" if cfg.apply_hg_interference_correction else "No"),
                 ("Pb-standard calibration requested", "Yes" if cfg.pb_standard_calibration.enabled else "No")]
    selected = set(_ratio_names(result))
    records = [r for s in result.samples for name, r in calibration_records(s).items() if name in selected]
    roles = Counter(next(iter(calibration_records(s).values())).role for s in result.samples if calibration_records(s))
    rows += [(f"Calibration role: {role}", n) for role, n in sorted(roles.items())]
    rows += [(f"Calibration ratio records: {status}", n) for status, n in sorted(Counter(r.status for r in records).items())]
    for mode in sorted({r.requested_mode for r in records}):
        rows.append(("Calibration mode", mode))
    seen = set()
    for record in records:
        reference_labels = {"material": "material", "value": "accepted ratio", "uncertainty": "uncertainty",
                            "k": "certificate k", "uncertainty_semantics": "uncertainty meaning"}
        for key, label in reference_labels.items():
            if key not in record.reference:
                continue
            value = record.reference[key]
            if isinstance(value, (str, int, float)) or value is None:
                pair = (f"Calibration reference {format_isotope_label(record.ratio_name)}: {label}", value)
                if pair not in seen:
                    rows.append(pair)
                    seen.add(pair)
    for s in result.samples:
        for record in hg_records(s).values():
            text = f"{record.status}: {record.source or 'source unavailable'}; {record.intensity_basis}; {record.reason}"
            pair = ("Hg correction evidence", text)
            if pair not in seen:
                rows.append(pair)
                seen.add(pair)
            for species, reference in (("Hg", record.hg_reference), ("Hg-support Tl", record.tl_reference)):
                for key in ("ratio_name", "value", "uncertainty", "k", "uncertainty_semantics"):
                    if key in reference:
                        pair = (f"{species} reference {key}", reference[key])
                        if pair not in seen:
                            rows.append(pair)
                            seen.add(pair)
    rows.append(("Drift applied", "Yes" if any(layer and layer.key == "drift" for layer, b, reason in contexts.values()) else "No"))
    if _delta_requested(result, cfg):
        rows.append(("Calibrated delta uncertainty", "Not calculated"))
    reasons = sorted({reason for layer, b, reason in contexts.values() if reason})
    rows += [("Result availability", reason) for reason in reasons]
    warnings = list(dict.fromkeys(result.warnings + [w for s in result.samples for w in s.warnings]))
    rows += [("Warning", w) for w in warnings]
    for s in result.samples:
        for b in s.uncertainty.values():
            for c in b.contributors:
                if c.state in {"MISSING_DATA", "NO_APPROVED_MODEL"}:
                    pair = ("Contributor status", f"{c.display_name or c.name}: {c.state}")
                    if pair not in seen:
                        rows.append(pair)
                        seen.add(pair)
            for item in b.coverage_limitations:
                note = item.get("note") or item.get("reason") or item.get("explanation")
                if note and ("Uncertainty limitation", note) not in seen:
                    rows.append(("Uncertainty limitation", note))
                    seen.add(("Uncertainty limitation", note))
    for row in rows:
        _append(ws, row)
    _finish(ws)
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 100
    for row in ws.iter_rows(min_row=2):
        ws.row_dimensions[row[0].row].height = max(32, 16 * math.ceil(len(str(row[1].value or "")) / 85))
        if isinstance(row[1].value, float):
            row[1].number_format = "0.###############"
    ws.freeze_panes = "B2"
    ws.print_title_rows = "1:1"


def export_pb_report(result, *, options=None, cycle_result=None, processing_config=None,
                     uncertainty_config=None, loaded_filename=None, element_config=None,
                     calibration_freshness=None, provenance=None, include_metadata_sheets=True):
    options = resolve_pb_options(options)
    validate_pb_report(result, processing_config, uncertainty_config, calibration_freshness)
    contexts = _contexts(result, processing_config)
    names = pb_sheet_names(options, has_detail=any(b is not None for layer, b, reason in contexts.values()))
    wb = Workbook()
    wb._traceiso_include_metadata_sheets = include_metadata_sheets
    _summary(wb, result, processing_config, uncertainty_config, loaded_filename, element_config, contexts)
    _results(wb, result, processing_config, uncertainty_config, contexts)
    if "Budget_Detail" in names:
        _detail(wb, result, uncertainty_config, contexts)
    if "Cycle_Data" in names:
        _cycle_data(wb, cycle_result if cycle_result is not None else result, processing_config)
    if wb.sheetnames != names:
        raise ValueError("Pb report sheet plan mismatch")
    output = BytesIO()
    from file_io.excel_writer import finalize_workbook_evidence
    finalize_workbook_evidence(wb, result, provenance=provenance, uncertainty_config=uncertainty_config)
    wb.save(output)
    output.seek(0)
    return output
