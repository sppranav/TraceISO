"""Compact spreadsheet export for the visible summary table."""

from __future__ import annotations

from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def build_summary_excel(df: pd.DataFrame) -> bytes:
    """Build a formatted Excel workbook from spreadsheet-safe summary data."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Summary"

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="Arial", size=10)

    for column_index, column_name in enumerate(df.columns, start=1):
        cell = worksheet.cell(row=1, column=column_index, value=column_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    for row_index, row in enumerate(df.itertuples(index=False, name=None), start=2):
        for column_index, value in enumerate(row, start=1):
            if pd.isna(value):
                value = None
            cell = worksheet.cell(row=row_index, column=column_index, value=value)
            cell.font = body_font
            cell.alignment = Alignment(vertical="center")

    worksheet.freeze_panes = "D2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.sheet_view.showGridLines = False
    worksheet.row_dimensions[1].height = 30

    for column_index, column_name in enumerate(df.columns, start=1):
        values = [str(column_name)]
        values.extend(str(value) for value in df.iloc[:, column_index - 1].dropna().head(200))
        width = min(max(len(value) for value in values) + 2, 32)
        worksheet.column_dimensions[get_column_letter(column_index)].width = max(width, 10)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
