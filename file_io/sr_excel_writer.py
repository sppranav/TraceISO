"""Scientific Sr report from recorded layers and aligned Engine-A budgets.

No correction is replayed here. Report rounding is confined to numeric output
cells; the supplied observations and canonical uncertainty fields stay intact.
"""

from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
from datetime import datetime, timezone
from io import BytesIO
import math
import textwrap

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from config.constants import APP_VERSION
from domain.layers import cycle_data_equal
from domain.models import ProcessingResult
from domain.ratio_selection import SR_STANDARD_LAYER_LABEL, select_best_ratio_layer
from domain.sr_standard_calibration import sr_calibration_record
from domain.uncertainty.scope import is_invalid_budget_scope
from file_io.report_formatting import report_uncertainty
from file_io.sanitize import format_isotope_label, literal_excel_text
from file_io.ssb_excel_writer import _header as _styled_header, _pair, _table

EXCEL_MAX_ROWS = 1_048_576
EXCEL_MAX_COLUMNS = 16_384


@dataclass(frozen=True)
class SrReportOptions:
    include_uncertainty_detail: bool = True
    include_cycle_data: bool = False


def resolve_sr_options(options=None, *, detail=True, cycles=False) -> SrReportOptions:
    if options is None:
        return SrReportOptions(detail, cycles)
    if not isinstance(options, SrReportOptions):
        raise TypeError("sr_report_options must be an SrReportOptions instance")
    return options


def sr_sheet_names(options: SrReportOptions, *, has_detail: bool) -> list[str]:
    options = resolve_sr_options(options)
    names = ["Summary", "Results"]
    if options.include_uncertainty_detail and has_detail:
        names.append("Budget_Detail")
    if options.include_cycle_data:
        names.append("Cycle_Data")
    return names


def _finite(value):
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def _header(ws, row, headers):
    _styled_header(ws, row, [literal_excel_text(h) if isinstance(h, str) else h for h in headers])


def _finish(ws):
    """Readable printed and screen tables, with precision set by cell meaning."""
    ws.sheet_view.showGridLines = False
    for col in range(1, ws.max_column + 1):
        width = max((len(str(ws.cell(row, col).value or "")) for row in range(1, min(ws.max_row, 80) + 1)), default=10)
        ws.column_dimensions[get_column_letter(col)].width = min(36, max(12, width + 2))
    current_headers = {}
    body_font = Font(name="Arial", size=11)
    alignments = {name: Alignment(wrap_text=True, vertical="center", horizontal=name, indent=1)
                  for name in ("left", "right", "center")}
    for row in ws:
        is_header = row[0].fill.fgColor.rgb == "002E75B6"
        if is_header:
            current_headers = {c.column: str(c.value) for c in row}
        height = 54 if is_header else 28
        for cell in row:
            heading = current_headers.get(cell.column, "")
            numeric = isinstance(cell.value, (int, float))
            if not is_header:
                cell.font = body_font
            cell.alignment = alignments["center" if is_header else "right" if numeric else "left"]
            if isinstance(cell.value, float) and cell.number_format == "General":
                magnitude = abs(cell.value)
                digits = 2 if ("(‰)" in heading or heading in {"u_c (ratio)", "U (ratio)"}) else 4
                places = max(0, digits - 1 - math.floor(math.log10(magnitude))) if magnitude else 0
                cell.number_format = "0." + "0" * places if places else "0"
            if cell.value is not None and not numeric:
                width = int(ws.column_dimensions[cell.column_letter].width - 2)
                lines = sum(max(1, len(textwrap.wrap(line, width=max(8, width)))) for line in str(cell.value).splitlines())
                height = max(height, 15 * lines + 8)
        ws.row_dimensions[row[0].row].height = min(409, height)
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = False
    ws.print_title_cols = "A:D" if ws.title in {"Results", "Cycle_Data"} else "A:B"


def _analyte(sample):
    return not sample.is_blank and not sample.is_standard


def _ratio_names(result):
    return sorted({name for sample in result.samples for attr in (
        "ratios", "blank_corrected_ratios", "interference_corrected_ratios", "corrected_ratios", "iif_corrected_ratios",
        "sr_standard_corrected_ratios", "drift_corrected_ratios", "uncertainty",
    ) for name in getattr(sample, attr)})


def _budget_context(sample, ratio):
    """Require the Engine-A final mean and support, never an earlier-stage U.

    Engine A records both ratio_value and basis_ratio_value on the final
    runtime layer (including committed calibration/drift). It has no separate
    persisted layer-name field; the canonical selector, mean and count jointly
    establish that contract without inventing provenance.
    """
    layer = select_best_ratio_layer(sample, ratio)
    budget = sample.uncertainty.get(ratio)
    if sr_calibration_record(sample, ratio).get("status") == "applied" and ratio not in sample.sr_standard_corrected_ratios:
        return None, None, "Recorded Sr calibration layer missing"
    if layer is None or layer.key not in {"iif", "sr_standard", "drift"}:
        return None, None, "Final Sr layer unavailable"
    if budget is None:
        return layer, None, "Budget unavailable"
    if is_invalid_budget_scope(budget):
        return layer, None, budget.scope_note or "Budget unavailable"
    if budget.engine != "internal_normalization" or budget.output_mode not in {"", "absolute_ratio"}:
        return layer, None, "Budget engine or measurand mismatch"
    mean = _finite(layer.data.mean)
    if mean is None or mean <= 0 or layer.data.n_valid < 2:
        return layer, None, "Insufficient accepted cycles"
    for basis in (budget.ratio_value, budget.basis_ratio_value):
        if _finite(basis) is None or not math.isclose(float(basis), mean, rel_tol=1e-12, abs_tol=0):
            return layer, None, "Budget does not describe the reported layer"
    if budget.n_cycles != layer.data.n_valid:
        return layer, None, "Budget accepted-cycle count mismatch"
    numbers = [budget.u_combined_abs, budget.expanded_abs,
               budget.u_combined_rel_permil, budget.expanded_rel_permil, budget.coverage_factor_k]
    if any(_finite(v) is None or v < 0 for v in numbers) or budget.coverage_factor_k <= 0:
        return layer, None, "Invalid uncertainty values"
    for actual, expected in (
        (budget.expanded_abs, budget.u_combined_abs * budget.coverage_factor_k),
        (budget.u_combined_rel_permil, budget.u_combined_abs / mean * 1000),
        (budget.expanded_rel_permil, budget.expanded_abs / mean * 1000),
    ):
        if not math.isclose(actual, expected, rel_tol=1e-10, abs_tol=0):
            return layer, None, "Inconsistent uncertainty units or coverage"
    return layer, budget, ""


def sr_has_detail(result: ProcessingResult) -> bool:
    return any(_budget_context(s, r)[1] is not None
               for s in result.samples if _analyte(s) for r in s.uncertainty)


def _stages(sample, ratio):
    """Retain scientific transitions even when their numerical factor is unity."""
    stages = {}
    blank = sample.blank_corrected_ratios.get(ratio)
    interference = sample.interference_corrected_ratios.get(ratio)
    corrected = sample.corrected_ratios.get(ratio)
    if blank is not None:
        stages["blank_corrected"] = ("Blank-corrected", blank)
    if interference is not None:
        stages["interference"] = ("Interference-corrected", interference)
    if corrected is not None and not cycle_data_equal(corrected, blank) and not cycle_data_equal(corrected, interference):
        inferred_interference = blank is not None and interference is None
        stages["interference" if inferred_interference else "corrected"] = (
            "Interference-corrected" if inferred_interference else "Corrected", corrected,
        )
    for key, label, mapping in (
        ("iif", "Internally normalized", sample.iif_corrected_ratios),
        ("sr_standard", SR_STANDARD_LAYER_LABEL, sample.sr_standard_corrected_ratios),
        ("drift", "Drift-corrected", sample.drift_corrected_ratios),
    ):
        if ratio in mapping:
            stages[key] = (label, mapping[ratio])
    return stages


def _stats(cd):
    return [_finite(cd.mean), _finite(cd.sd), _finite(cd.se)] if cd is not None else [None] * 3


def _coverage_heading():
    return CellRichText(TextBlock(InlineFont(b=True, i=True, color="FFFFFFFF"), "k"))


def _uncertainty_headers():
    return ["u_c (ratio)", "U (ratio)", "Relative ratio u_c (‰)", "Relative ratio U (‰)",
            _coverage_heading(), "Effective degrees of freedom", "Dominant contributor"]


def _uncertainty_values(budget, *, show_dof=True):
    if budget is None:
        return [None] * 7
    dof = budget.effective_dof
    return [report_uncertainty(v) for v in (
        budget.u_combined_abs, budget.expanded_abs,
        budget.u_combined_rel_permil, budget.expanded_rel_permil,
    )] + [budget.coverage_factor_k, ("∞" if dof == math.inf else _finite(dof)) if show_dof else None, budget.dominant_contributor]


def _append(ws, values):
    ws.append([literal_excel_text(v) if isinstance(v, str) else v for v in values])


def _results(wb, result, uncertainty_config=None):
    ws = wb.create_sheet("Results")
    observations = [s for s in result.samples if not s.is_blank]
    for ratio in _ratio_names(result):
        per_sample = [_stages(s, ratio) for s in observations]
        schema = {}
        for key in ("blank_corrected", "interference", "corrected", "iif", "sr_standard", "drift"):
            for stages in per_sample:
                if key in stages:
                    schema[key] = stages[key][0]
                    break
        _append(ws, [format_isotope_label(ratio)])
        header_row = ws.max_row + 1
        headers = ["Sample", "Type", "Run", "Accepted n (reported)"]
        counts = {}
        for key, label in schema.items():
            counts[key] = any(
                key in stages and (select_best_ratio_layer(s, ratio) is None or
                stages[key][1].n_valid != select_best_ratio_layer(s, ratio).data.n_valid)
                for s, stages in zip(observations, per_sample)
            )
            headers += [f"{label} ratio", f"{label} SD", f"{label} SE"]
            if counts[key]:
                headers.append(f"{label} n")
        value_col = len(headers) + 1
        headers += ["Reported ratio"] + _uncertainty_headers() + ["Uncertainty availability"]
        _header(ws, header_row, headers)
        for sample, stages in zip(observations, per_sample):
            layer, budget, reason = _budget_context(sample, ratio)
            values = [sample.name, sample.sample_type, sample.run_number,
                      layer.data.n_valid if layer is not None else None]
            for key in schema:
                cd = stages[key][1] if key in stages else None
                values += _stats(cd)
                if counts[key]:
                    values.append(cd.n_valid if cd is not None else None)
            values += [_finite(layer.data.mean) if layer is not None else None]
            values += _uncertainty_values(budget, show_dof=getattr(uncertainty_config, "coverage_method", "") != "fixed_k")
            values += [f"Not quantified: {reason}" if budget is None and _analyte(sample) else None]
            _append(ws, values)
            for col, heading in enumerate(headers, 1):
                if str(heading).endswith(" ratio"):
                    ws.cell(ws.max_row, col).number_format = "0.00000000"
            _pair(ws.cell(ws.max_row, value_col), ws.cell(ws.max_row, value_col + 2))
        _table(ws, header_row, ws.max_row, headers)
        ws.append([])
    ws.freeze_panes = "E3"
    ws.print_title_rows = "1:2"
    _finish(ws)


_MISSING_STATES = {"MISSING_DATA", "NO_APPROVED_MODEL"}


def _contributor_value(contributor):
    if contributor.state in _MISSING_STATES:
        return "Not quantified"
    if contributor.state != "ACTIVE":
        return None
    value = _finite(contributor.value_rel_permil)
    return report_uncertainty(value) if value is not None and value >= 0 else "Not quantified"


def _budget_detail(wb, result, uncertainty_config=None):
    ws = wb.create_sheet("Budget_Detail")
    for ratio in _ratio_names(result):
        rows = [(s, *_budget_context(s, ratio)) for s in result.samples if _analyte(s)]
        columns = {}
        applicable = set()
        for sample, _layer, budget, _reason in rows:
            source = budget or sample.uncertainty.get(ratio)
            if source is None or source.engine != "internal_normalization":
                continue
            for c in source.contributors:
                if c.state == "ACTIVE" or c.state in _MISSING_STATES:
                    applicable.add(c.name)
                label = c.display_name or c.name
                if c.name.startswith("u_std_repeatability"):
                    label += " (SRM-session basis)"
                columns.setdefault(c.name, f"{label} [{c.name}] (‰)")
        columns = {name: label for name, label in columns.items() if name in applicable}
        _append(ws, [format_isotope_label(ratio)])
        header_row = ws.max_row + 1
        headers = ["Sample", "Run", "Reported source layer", "Budget-basis ratio", "Cycle SE (ratio)"]
        headers += list(columns.values()) + _uncertainty_headers() + ["Uncertainty availability"]
        _header(ws, header_row, headers)
        for sample, layer, budget, reason in rows:
            contributors = {c.name: c for c in budget.contributors} if budget else {}
            values = [sample.name, sample.run_number, layer.label if layer else None,
                      budget.basis_ratio_value if budget else None,
                      _finite(layer.data.se) if layer is not None and budget else None]
            values += [_contributor_value(contributors[name]) if name in contributors else None for name in columns]
            values += _uncertainty_values(budget, show_dof=getattr(uncertainty_config, "coverage_method", "") != "fixed_k")
            values += [f"Not quantified: {reason}" if budget is None else None]
            _append(ws, values)
            ws.cell(ws.max_row, 4).number_format = "0.00000000"
            _pair(ws.cell(ws.max_row, 4), ws.cell(ws.max_row, 7 + len(columns)))
        _table(ws, header_row, ws.max_row, headers)
        ws.append([])
    ws.freeze_panes = "C3"
    ws.print_title_rows = "1:2"
    _finish(ws)


def _series(sample):
    series = []
    for key, label, mapping in (
        (0, "Raw intensity", sample.intensities),
        (1, "Blank-corrected intensity", sample.blank_corrected_intensities),
        (2, "Interference-corrected intensity", sample.interference_corrected_intensities),
        (3, "Corrected intensity", sample.corrected_intensities),
        (4, "Raw ratio", sample.ratios),
    ):
        for name, cd in sorted(mapping.items()):
            if key == 3 and any(n == name and cycle_data_equal(cd, prior)
                                for k, _label, n, prior in series if k in {1, 2}):
                continue
            # Sr stores final interference intensities in corrected_intensities.
            stage = ("Interference-corrected intensity" if key == 3 and sample.metadata.get("_sr_chain_method")
                     and name not in sample.interference_corrected_intensities else label)
            rank = 2 if stage == "Interference-corrected intensity" else key
            series.append((rank, stage, name, cd))
    for ratio in sorted({r for attr in ("blank_corrected_ratios", "interference_corrected_ratios", "corrected_ratios", "iif_corrected_ratios",
                                        "sr_standard_corrected_ratios", "drift_corrected_ratios")
                         for r in getattr(sample, attr)}):
        for key, (label, cd) in _stages(sample, ratio).items():
            rank = {"blank_corrected": 5, "interference": 6, "corrected": 7, "iif": 8, "sr_standard": 9, "drift": 10}[key]
            series.append((rank, f"{label} ratio", ratio, cd))
    return series


def _cycle_data(wb, result):
    prepared = [(s, _series(s)) for s in result.samples]
    specifications = sorted({(rank, label, name) for _s, series in prepared for rank, label, name, _cd in series})
    counts = [max((cd.n_total for *_key, cd in series), default=0) for _s, series in prepared]
    rows, columns = 1 + sum(counts), 4 + 2 * len(specifications)
    if rows > EXCEL_MAX_ROWS or columns > EXCEL_MAX_COLUMNS:
        raise ValueError(f"Cycle_Data dimensions {rows} rows x {columns} columns exceed Excel limits ({EXCEL_MAX_ROWS} x {EXCEL_MAX_COLUMNS})")
    ws = wb.create_sheet("Cycle_Data")
    headers = ["Sample", "Type", "Run", "Cycle"]
    for rank, label, name in specifications:
        heading = f"{label}: {format_isotope_label(name)}" + (" (V)" if rank < 4 else "")
        headers += [heading, heading + " status"]
    _header(ws, 1, headers)
    duplicates = Counter((s.name, s.run_number) for s in result.samples)
    row_number = 2
    for (sample, series), count in zip(prepared, counts):
        lookup = {(rank, label, name): cd for rank, label, name, cd in series}
        name = sample.name
        if duplicates[(name, sample.run_number)] > 1:
            name += f" [{sample.observation_id}]"
        for index in range(count):
            values = [name, sample.sample_type, sample.run_number, index + 1]
            for key in specifications:
                cd = lookup.get(key)
                if cd is None or index >= cd.n_total:
                    values += [None, "Not recorded"]
                else:
                    value = _finite(cd.values[index])
                    values += [value if value is not None else str(cd.values[index]),
                               "Accepted" if cd.mask[index] else "Excluded"]
            _append(ws, values)
            for offset, (rank, _label, _name) in enumerate(specifications):
                ws.cell(row_number, 5 + 2 * offset).number_format = "0.000000E+00" if rank < 4 else "0.0000000000"
            row_number += 1
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:{get_column_letter(columns)}{rows}"
    ws.print_title_rows = "1:1"
    _finish(ws)


def _summary(wb, result, processing_config, uncertainty_config, loaded_filename, element_config):
    ws = wb.active
    ws.title = "Summary"
    _header(ws, 1, ["Sr scientific report", "Session context"])
    rows = [("Source", loaded_filename), ("Software", f"TraceISO {APP_VERSION}"),
            ("Export UTC", datetime.now(timezone.utc).isoformat()),
            ("Workflow", "Sr internal normalization"),
            ("Reported ratios", ", ".join(format_isotope_label(r) for r in _ratio_names(result)))]
    methods = sorted({s.metadata.get("_sr_chain_method") for s in result.samples if s.metadata.get("_sr_chain_method")})
    labels = {"sr_natural_init_two_refinements_v1": "Natural-ratio initialization and two refinements",
              "sr_single_pass_measured_f_v1": "Single-pass Sr correction"}
    for method in methods:
        rows.append(("Sr method", f"{labels.get(method, method)} ({method})"))
    cfg, uc = processing_config, uncertainty_config
    if cfg is not None:
        norm_ratio = cfg.normalization_ratio_override or getattr(element_config, "normalization_ratio", None)
        norm_value = cfg.normalization_value_override
        if norm_value is None:
            norm_value = getattr(element_config, "normalization_value", None)
        rows += [("Normalization ratio", format_isotope_label(norm_ratio) if norm_ratio else None),
                 ("Accepted normalization value", _finite(norm_value)),
                 ("Normalization override", "Yes" if cfg.normalization_value_override is not None or cfg.normalization_ratio_override else "No"),
                 ("Blank method", cfg.blank_mode),
                 ("Sr 3-variable blank model", "Yes (⁸⁷Sr/⁸⁶Sr)" if getattr(uc, "sr_blank_3var", False) and cfg.blank_mode != "none" else "No"),
                 ("Outlier method", cfg.filter_method)]
        if cfg.filter_method != "None":
            rows.append(("Outlier threshold", cfg.get_active_filter_threshold()))
        roles = getattr(element_config, "correction_roles", {})
        for term, role, monitor in (("Rb", "rb_interfering_mass", "rb_monitor"),
                                    ("⁸⁶Kr", "kr86_interfering_mass", "kr_monitor"),
                                    ("⁸⁴Kr", "kr84_interfering_mass", "kr_monitor")):
            mass = roles.get(role)
            enabled = cfg.apply_interference_correction and mass is not None and cfg.is_monitor_enabled(mass)
            rows.append((f"{term} interference correction", f"Enabled ({format_isotope_label(roles.get(monitor, 'monitor unassigned'))})" if enabled else "Disabled"))
        rows.append(("Sr-standard calibration requested", "Yes" if cfg.sr_session_anchoring else "No"))
    records = [sr_calibration_record(s) for s in result.samples if sr_calibration_record(s)]
    applied = next((r for r in records if r.get("status") == "applied"), None)
    rows.append(("Sr-standard calibration applied", "Yes" if applied else "No"))
    if applied:
        rows += [("Calibration mode", applied.get("method")),
                 ("Calibration reference", getattr(cfg, "reference_material", None)),
                 ("Included calibration standards", len(applied.get("standards", [])))]
    drift = any(s.drift_corrected_ratios for s in result.samples)
    rows.append(("Drift applied", "Yes" if drift else "No"))
    if drift and cfg is not None:
        rows += [("Drift method", cfg.drift.method), ("Drift degree", cfg.drift.degree)]
    rows += [(f"Observations: {kind}", count) for kind, count in sorted(Counter(s.sample_type for s in result.samples).items())]
    if uc is not None:
        rows.append(("Coverage method", {"fixed_k": "Fixed coverage factor", "welch_satterthwaite": "Welch–Satterthwaite"}.get(uc.coverage_method, uc.coverage_method)))
    factors = {b.coverage_factor_k for s in result.samples for r in s.uncertainty
               if (b := _budget_context(s, r)[1]) is not None}
    if len(factors) == 1:
        rows.append(("Coverage k", factors.pop()))
    rows.append(("Cross-ratio scope", "No cross-ratio covariance is exported. Do not infer independence between reported ratios."))
    warnings = list(dict.fromkeys(result.warnings + [w for s in result.samples for w in s.warnings]))
    warnings = list(dict.fromkeys(
        "Sr refinement-stability warning recorded. The fixed two-refinement result retains a truncation limitation."
        if "refinement" in w.lower() else w for w in warnings
    ))
    rows += [("Warning", warning) for warning in warnings]
    limitations = set()
    for sample in result.samples:
        for budget in sample.uncertainty.values():
            if budget.engine != "internal_normalization":
                continue
            for c in budget.contributors:
                if c.state in _MISSING_STATES:
                    limitations.add(f"{c.display_name or c.name}: Not quantified")
            for limitation in budget.coverage_limitations:
                note = limitation.get("note") or limitation.get("reason")
                if note:
                    limitations.add(note)
    rows += [("Uncertainty limitation", note) for note in sorted(limitations)]
    for row in rows:
        _append(ws, row)
        if row[0] == "Coverage k":
            ws.cell(ws.max_row, 1).value = CellRichText(
                "Coverage ", TextBlock(InlineFont(i=True), "k"),
            )
    _finish(ws)
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 100
    for row in range(2, ws.max_row + 1):
        ws.row_dimensions[row].height = max(32, 16 * math.ceil(len(str(ws.cell(row, 2).value or "")) / 90))
    ws.freeze_panes = "B2"
    ws.print_title_rows = "1:1"


def export_sr_report(result: ProcessingResult, *, options: SrReportOptions,
                     cycle_result=None, processing_config=None, uncertainty_config=None,
                     loaded_filename=None, element_config=None, provenance=None, include_metadata_sheets=True) -> BytesIO:
    options = resolve_sr_options(options)
    if result.element_symbol != "Sr":
        raise ValueError("The Sr report requires an Sr session")
    if processing_config is not None:
        if processing_config.sr_session_anchoring and not any(sr_calibration_record(s) for s in result.samples if not s.is_blank):
            raise ValueError("Sr export refused: requested calibration has no applied-state record; reprocess the session")
        drift_cfg = processing_config.drift
        if drift_cfg.enabled and not any(drift_cfg.ratio_name in s.drift_corrected_ratios for s in result.samples):
            raise ValueError("Sr export refused: requested drift correction is unavailable; reprocess the session")
        if drift_cfg.enabled and any(
            not s.is_blank and drift_cfg.ratio_name in s.iif_corrected_ratios
            and drift_cfg.ratio_name not in s.drift_corrected_ratios for s in result.samples
        ):
            raise ValueError("Sr export refused: requested drift correction is missing for one or more observations")
    wb = Workbook()
    wb._traceiso_include_metadata_sheets = include_metadata_sheets
    _summary(wb, result, processing_config, uncertainty_config, loaded_filename, element_config)
    _results(wb, result, uncertainty_config)
    names = sr_sheet_names(options, has_detail=sr_has_detail(result))
    if "Budget_Detail" in names:
        _budget_detail(wb, result, uncertainty_config)
    if "Cycle_Data" in names:
        _cycle_data(wb, cycle_result if cycle_result is not None else result)
    if wb.sheetnames != names:
        raise ValueError("Sr workbook sheet plan mismatch")
    output = BytesIO()
    from file_io.excel_writer import finalize_workbook_evidence
    finalize_workbook_evidence(wb, result, provenance=provenance, uncertainty_config=uncertainty_config)
    wb.save(output)
    output.seek(0)
    return output
